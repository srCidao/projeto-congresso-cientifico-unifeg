"""
Corta Frutas - Explosão
-----------------------
Jogo interativo desenvolvido em Python com pygame e openpyxl. 
O objetivo é cortar as frutas lançadas pela tela com o cursor, acumulando pontos dentro de um tempo limite.  
A estrutura do sistema inclui três partes principais:
1. Cadastro do jogador, com armazenamento em planilha Excel.
2. Sessão de jogo com geração e corte de frutas animadas.
3. Exibição de ranking com as maiores pontuações registradas.

O fluxo segue a sequência: Cadastro → Jogo → Ranking, permitindo recomeçar novas partidas.  
O pygame é utilizado para renderização gráfica, controle de eventos e animações, 
enquanto o openpyxl gerencia os dados de pontuação e cadastro no arquivo Excel.  
As frutas são lançadas da parte inferior da tela e cortadas quando o cursor colide com elas, 
gerando explosões de partículas e somando pontos ao jogador.  
Ao final do tempo, o placar é salvo e exibido no ranking geral.
"""

import pygame
import random
import math
import sys
import ctypes
import openpyxl
import os
from collections import defaultdict

# ============================
# CONFIGURAÇÕES / CONSTANTES
# ============================
FPS = 60
TEMPO_JOGO = 20
VELOCIDADE_FRUTAS = 0.1
PONTOS_POR_FRUTA = 10

# Paleta de cores RGB
BRANCO = (255, 255, 255)
PRETO = (0, 0, 0)
VERMELHO = (255, 0, 0)
VERDE = (0, 255, 0)
AZUL = (0, 120, 255)
AMARELO = (255, 255, 0)
LARANJA = (255, 165, 0)
ROXO = (128, 0, 128)
FUNDO = (245, 250, 255)

CORES_FRUTA = [VERMELHO, LARANJA, AMARELO, ROXO, VERDE, AZUL]
ARQUIVO_EXCEL = "jogadores.xlsx"

# ============================
# INICIALIZAÇÃO PYGAME / TELA
# ============================
pygame.init()

# Criação da janela principal (redimensionável)
tela = pygame.display.set_mode((800, 600), pygame.RESIZABLE)

# Maximiza a janela no Windows
try:
    user32 = ctypes.windll.user32
    hwnd = pygame.display.get_wm_info().get('window')
    if hwnd:
        user32.ShowWindow(hwnd, 3)  # SW_MAXIMIZE = 3
except Exception:
    pass

LARGURA, ALTURA = tela.get_size()
pygame.display.set_caption("Corta Frutas - Explosão")
relogio = pygame.time.Clock()

# Carrega o cursor personalizado
try:
    cursor_img = pygame.image.load("cursor_mao.png").convert_alpha()
    cursor_img = pygame.transform.scale(cursor_img, (64, 64))
    pygame.mouse.set_visible(False)
except Exception:
    print("⚠️  Arquivo 'cursor_mao.png' não encontrado!")
    pygame.quit()
    sys.exit()

# Fontes pré-carregadas para melhor desempenho
FONT_SMALL = pygame.font.SysFont(None, 24)
FONT_MED = pygame.font.SysFont(None, 32)
FONT_LARGE = pygame.font.SysFont(None, 48)
FONT_HUGE = pygame.font.SysFont(None, 72)

# ============================
# FUNÇÕES DE I/O COM EXCEL
# ============================

def salvar_dados_excel(nome, idade, telefone, email, obs, pontuacao):
    """
    Registra uma linha no arquivo Excel com os dados do jogador e pontuação.
    Cria o arquivo e cabeçalho automaticamente se ainda não existir.
    """
    arquivo = ARQUIVO_EXCEL
    if not os.path.exists(arquivo):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Pontuações"
        ws.append(["Nome", "Idade", "Telefone", "Email", "Observações", "Pontuação"])
        wb.save(arquivo)
        wb.close()
    wb = openpyxl.load_workbook(arquivo)
    ws = wb.active
    try:
        p_val = int(pontuacao)
    except Exception:
        p_val = pontuacao
    ws.append([nome, idade, telefone, email, obs, p_val])
    wb.save(arquivo)
    wb.close()


def ler_ranking_excel(top=10):
    """
    Retorna os N melhores jogadores com base na pontuação.
    Cada item do retorno é um dicionário com nome e pontuação.
    """
    arquivo = ARQUIVO_EXCEL
    if not os.path.exists(arquivo):
        return []
    wb = openpyxl.load_workbook(arquivo, data_only=True)
    ws = wb.active
    dados = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        nome = row[0]
        pont = row[5] if len(row) > 5 else None
        try:
            pont = int(pont) if pont is not None else 0
        except Exception:
            pont = 0
        if nome is None:
            continue
        dados.append({"nome": str(nome), "pontuacao": pont})
    wb.close()
    dados = sorted(dados, key=lambda x: x["pontuacao"], reverse=True)
    return dados[:top]


def ler_jogadores_unicos():
    """
    Retorna a melhor pontuação registrada de cada jogador único.
    """
    arquivo = ARQUIVO_EXCEL
    if not os.path.exists(arquivo):
        return []
    wb = openpyxl.load_workbook(arquivo, data_only=True)
    ws = wb.active
    pontuacoes = defaultdict(int)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        nome = row[0]
        pont = row[5] if len(row) > 5 else None
        if nome is None:
            continue
        try:
            pont = int(pont) if pont is not None else 0
        except Exception:
            pont = 0
        if pont > pontuacoes[nome]:
            pontuacoes[nome] = pont
    wb.close()
    return [{"nome": n, "pontuacao": p} for n, p in pontuacoes.items()]

# ============================
# CLASSES: Particula e Fruta
# ============================

class Particula:
    """Representa partículas geradas ao cortar uma fruta."""
    def __init__(self, x, y, cor, tamanho=5, velocidade=6):
        self.x, self.y = x, y
        self.vel_x = random.uniform(-velocidade, velocidade)
        self.vel_y = random.uniform(-velocidade - 2, -1)
        self.raio = random.randint(tamanho // 2, tamanho)
        self.vida = random.randint(25, 50)
        self.alpha = 255
        self.cor_original = cor
        self.cor = cor

    def atualizar(self):
        """Atualiza movimento, gravidade e transparência da partícula."""
        self.x += self.vel_x
        self.y += self.vel_y
        self.vel_y += 0.4
        self.vida -= 1
        self.alpha = max(0, self.alpha - 6)
        fator = self.alpha / 255
        self.cor = tuple(max(0, min(255, int(c * fator))) for c in self.cor_original)
        self.raio = max(0, self.raio - 0.05)

    def desenhar(self, surf):
        """Renderiza a partícula na tela."""
        if self.alpha > 0 and self.raio > 0:
            superficie = pygame.Surface((int(self.raio * 2) + 2, int(self.raio * 2) + 2), pygame.SRCALPHA)
            pygame.draw.circle(superficie, (*self.cor, int(self.alpha)), (int(self.raio), int(self.raio)), int(max(1, self.raio)))
            surf.blit(superficie, (self.x - self.raio, self.y - self.raio))

    def ativa(self):
        """Retorna True se a partícula ainda está visível e ativa."""
        return self.vida > 0 and self.alpha > 0


class Fruta:
    """Controla o comportamento individual de cada fruta lançada na tela."""
    def __init__(self):
        self.cor = random.choice(CORES_FRUTA)
        self.tamanho = random.randint(40, 60)
        MARGEM_DIREITA = 400
        max_x = max(50, LARGURA - MARGEM_DIREITA - 50)
        if max_x <= 50:
            max_x = max(50, LARGURA - 50)
        self.x = random.randint(50, max_x)
        self.y = ALTURA + self.tamanho
        self.velocidade_x = random.uniform(-VELOCIDADE_FRUTAS * 1.5, VELOCIDADE_FRUTAS * 1.5)
        try:
            self.velocidade_y = -math.sqrt(max(0.1, 2 * 0.5 * (self.y - self.tamanho)))
        except Exception:
            self.velocidade_y = -10
        self.fatiada = False
        self.tempo_fatiada = 0

    def atualizar(self):
        """Atualiza a posição da fruta e seu estado."""
        self.velocidade_y += 0.6
        self.x += self.velocidade_x
        self.y += self.velocidade_y
        if self.fatiada:
            self.tempo_fatiada += 1

    def desenhar(self, surf):
        """Renderiza a fruta ou suas metades cortadas."""
        if not self.fatiada:
            pygame.draw.circle(surf, self.cor, (int(self.x), int(self.y)), self.tamanho)
        else:
            for offset in (-15, 15):
                pygame.draw.circle(surf, self.cor, (int(self.x + offset), int(self.y)), self.tamanho // 2)

    def esta_fora_da_tela(self):
        """Retorna True se a fruta saiu da tela ou já foi exibida cortada."""
        return self.y > ALTURA + self.tamanho or self.tempo_fatiada > 30

    def foi_fatiada(self, mouse_x, mouse_y):
        """Detecta se o cursor colidiu com a fruta."""
        return (not self.fatiada) and (math.hypot(self.x - mouse_x, self.y - mouse_y) < self.tamanho)

# ============================
# FUNÇÕES DE INTERFACE GRÁFICA
# ============================

def desenhar_texto(texto, tamanho, cor, x, y, centralizado=True):
    """Desenha texto centralizado ou alinhado, usando fontes pré-carregadas."""
    if tamanho >= 72:
        fonte = FONT_HUGE
    elif tamanho >= 48:
        fonte = FONT_LARGE
    elif tamanho >= 32:
        fonte = FONT_MED
    else:
        fonte = FONT_SMALL
    imagem = fonte.render(str(texto), True, cor)
    rect = imagem.get_rect(center=(x, y)) if centralizado else imagem.get_rect(topleft=(x, y))
    tela.blit(imagem, rect)


def desenhar_botao_texto(texto, rect, cor, texto_cor, border_radius=20, fonte_size=48):
    """Desenha um botão retangular com texto centralizado."""
    pygame.draw.rect(tela, cor, rect, border_radius=border_radius)
    desenhar_texto(texto, fonte_size, texto_cor, rect.centerx, rect.centery)


def criar_explosao(x, y, cor):
    """Gera uma lista de partículas de explosão em torno das coordenadas indicadas."""
    return [Particula(x, y, cor, tamanho=8, velocidade=7) for _ in range(40)]

# ============================
# LOOP PRINCIPAL DO JOGO
# ============================

def main():
    """
    Controla o ciclo completo do jogo: cadastro, partida e exibição de ranking.
    Gerencia eventos, atualizações de tela e fluxo entre as fases.
    """
    frutas = []
    particulas = []
    tempo_restante = TEMPO_JOGO
    pontuacao = 0
    jogo_ativo = False
    cadastro_finalizado = False
    mostrar_ranking = False
    selecionando_jogador = False
    ultimo_tempo_fruta = 0
    scroll_offset = 0
    dados = {"nome": "", "idade": "", "telefone": "", "email": "", "obs": ""}
    campos = list(dados.keys())
    campo_atual = 0
    ranking = []
    rodando = True
    mouse_clicado = False

    # Loop principal de execução
    while rodando:
        tela.fill(FUNDO)
        mouse_x, mouse_y = pygame.mouse.get_pos()
        mouse_clicado = False
        eventos = pygame.event.get()

        # Captura de eventos globais
        for evento in eventos:
            if evento.type == pygame.QUIT:
                rodando = False
            elif evento.type == pygame.KEYDOWN:
                if evento.key == pygame.K_ESCAPE:
                    rodando = False
                elif not cadastro_finalizado and not selecionando_jogador:
                    if evento.key in (pygame.K_TAB, pygame.K_DOWN, pygame.K_RETURN):
                        campo_atual = (campo_atual + 1) % len(campos)
                    elif evento.key == pygame.K_UP:
                        campo_atual = (campo_atual - 1) % len(campos)
                    elif evento.key == pygame.K_BACKSPACE:
                        dados[campos[campo_atual]] = dados[campos[campo_atual]][:-1]
                    else:
                        dados[campos[campo_atual]] += evento.unicode
            elif evento.type == pygame.MOUSEWHEEL and selecionando_jogador:
                scroll_offset -= evento.y * 40
            elif evento.type == pygame.MOUSEBUTTONDOWN and evento.button == 1:
                mouse_clicado = True

        # TELA DE SELEÇÃO DE JOGADOR EXISTENTE
        if selecionando_jogador:
            desenhar_texto("Selecionar Jogador", 48, PRETO, LARGURA // 2, 60)
            jogadores = ler_jogadores_unicos()
            total = len(jogadores)
            visiveis = 10
            max_scroll = max(0, (total - visiveis) * 50)
            scroll_offset = max(0, min(scroll_offset, max_scroll))
            inicio = int(scroll_offset // 50)
            fim = min(inicio + visiveis, total)
            base_y = 140

            for i, j in enumerate(jogadores[inicio:fim]):
                y = base_y + i * 50
                desenhar_texto(f"{j['nome']}  -  {j['pontuacao']} pts", 32, PRETO, LARGURA // 2 - 60, y, centralizado=False)
                botao_sel = pygame.Rect(LARGURA // 2 + 200, y - 20, 120, 40)
                cor_sel = AZUL if botao_sel.collidepoint(mouse_x, mouse_y) else ROXO
                pygame.draw.rect(tela, cor_sel, botao_sel, border_radius=10)
                desenhar_texto("Selecionar", 24, BRANCO, botao_sel.centerx, botao_sel.centery)
                if mouse_clicado and botao_sel.collidepoint(mouse_x, mouse_y):
                    dados["nome"] = j["nome"]
                    cadastro_finalizado = True
                    selecionando_jogador = False
                    scroll_offset = 0
                    break

            if total > visiveis:
                botao_up = pygame.Rect(LARGURA - 80, 150, 50, 50)
                botao_down = pygame.Rect(LARGURA - 80, ALTURA - 100, 50, 50)
                pygame.draw.rect(tela, AZUL if botao_up.collidepoint(mouse_x, mouse_y) else ROXO, botao_up, border_radius=10)
                pygame.draw.rect(tela, AZUL if botao_down.collidepoint(mouse_x, mouse_y) else ROXO, botao_down, border_radius=10)
                desenhar_texto("↑", 32, BRANCO, botao_up.centerx, botao_up.centery)
                desenhar_texto("↓", 32, BRANCO, botao_down.centerx, botao_down.centery)
                if mouse_clicado:
                    if botao_up.collidepoint(mouse_x, mouse_y):
                        scroll_offset = max(0, scroll_offset - 50)
                    elif botao_down.collidepoint(mouse_x, mouse_y):
                        scroll_offset = min(max_scroll, scroll_offset + 50)

        # TELA DE CADASTRO
        elif not cadastro_finalizado:
            desenhar_texto("Cadastro do Jogador", 48, PRETO, LARGURA // 2, 60)
            for i, campo in enumerate(campos):
                cor = AZUL if i == campo_atual else PRETO
                texto = f"{campo.capitalize()}: {dados[campo]}"
                desenhar_texto(texto, 32, cor, LARGURA // 2, 150 + i * 50)
            botao_confirmar = pygame.Rect(LARGURA // 2 + 50, ALTURA - 100, 200, 60)
            botao_reusar = pygame.Rect(LARGURA // 2 - 250, ALTURA - 100, 200, 60)
            cor_confirmar = AZUL if botao_confirmar.collidepoint(mouse_x, mouse_y) else ROXO
            cor_reusar = AZUL if botao_reusar.collidepoint(mouse_x, mouse_y) else ROXO
            pygame.draw.rect(tela, cor_confirmar, botao_confirmar, border_radius=20)
            pygame.draw.rect(tela, cor_reusar, botao_reusar, border_radius=20)
            desenhar_texto("CONFIRMAR", 32, BRANCO, botao_confirmar.centerx, botao_confirmar.centery)
            desenhar_texto("JOGAR NOVAMENTE", 26, BRANCO, botao_reusar.centerx, botao_reusar.centery)
            desenhar_texto("Preencha Nome, Idade e Telefone (obrigatórios)", 24, VERMELHO, LARGURA // 2, botao_confirmar.top - 30)

            # Detecta clique nos botões de confirmar ou selecionar jogador existente
            if mouse_clicado:
                if botao_confirmar.collidepoint(mouse_x, mouse_y):
                    if all(dados[k].strip() for k in ("nome", "idade", "telefone")):
                        cadastro_finalizado = True
                elif botao_reusar.collidepoint(mouse_x, mouse_y):
                    selecionando_jogador = True
                    campo_atual = 0  # Reinicia seleção

        # TELA DE PLAY (pré-jogo)
        elif (not jogo_ativo) and (not mostrar_ranking):
            botao_play = pygame.Rect(LARGURA // 2 - 100, ALTURA // 2 - 60, 200, 120)
            cor_play = AZUL if botao_play.collidepoint(mouse_x, mouse_y) else ROXO
            pygame.draw.rect(tela, cor_play, botao_play, border_radius=20)
            desenhar_texto("PLAY", 72, BRANCO, botao_play.centerx, botao_play.centery)

            # Inicia partida ao passar o cursor sobre o botão
            if botao_play.collidepoint(mouse_x, mouse_y):
                jogo_ativo = True
                tempo_restante = TEMPO_JOGO
                frutas.clear()
                particulas.clear()
                pontuacao = 0
                ultimo_tempo_fruta = 0

        # TELA DE RANKING
        elif mostrar_ranking:
            desenhar_texto("RANKING - TOP 10", 48, PRETO, LARGURA // 2, 60)
            for i, p in enumerate(ranking):
                texto = f"{i+1:>2}. {p['nome']} - {p['pontuacao']} pts"
                desenhar_texto(texto, 32, AZUL if i < 3 else PRETO, LARGURA // 2, 140 + i * 40)

            botao_voltar = pygame.Rect(LARGURA // 2 - 100, ALTURA - 120, 200, 60)
            cor_voltar = AZUL if botao_voltar.collidepoint(mouse_x, mouse_y) else ROXO
            pygame.draw.rect(tela, cor_voltar, botao_voltar, border_radius=20)
            desenhar_texto("JOGAR NOVAMENTE", 26, BRANCO, botao_voltar.centerx, botao_voltar.centery)

            # Reinicia cadastro e fluxo do jogo ao clicar no botão
            if mouse_clicado and botao_voltar.collidepoint(mouse_x, mouse_y):
                cadastro_finalizado = False
                mostrar_ranking = False
                dados = {k: "" for k in dados}
                campo_atual = 0
                ranking = []

        # LOOP DO JOGO (frutas, partículas e pontuação)
        elif jogo_ativo:
            # Atualiza tempo restante
            tempo_restante -= 1 / FPS
            if tempo_restante <= 0:
                # Fim de partida: salva dados e exibe ranking
                jogo_ativo = False
                try:
                    salvar_dados_excel(
                        dados["nome"], dados["idade"], dados["telefone"],
                        dados["email"], dados["obs"], pontuacao
                    )
                except Exception as e:
                    print("Erro ao salvar Excel:", e)
                ranking = ler_ranking_excel()
                mostrar_ranking = True
                continue  # Próximo frame mostra ranking

            tempo_atual = pygame.time.get_ticks()
            # Gera nova fruta a cada ~1000ms
            if tempo_atual - ultimo_tempo_fruta > 1000:
                frutas.append(Fruta())
                ultimo_tempo_fruta = tempo_atual

            # Atualiza frutas e verifica corte
            for fruta in frutas[:]:
                fruta.atualizar()
                if fruta.esta_fora_da_tela():
                    try:
                        frutas.remove(fruta)
                    except ValueError:
                        pass
                elif fruta.foi_fatiada(mouse_x, mouse_y):
                    fruta.fatiada = True
                    pontuacao += PONTOS_POR_FRUTA
                    particulas.extend(criar_explosao(fruta.x, fruta.y, fruta.cor))

            # Atualiza partículas e filtra inativas
            for p in particulas:
                p.atualizar()
            particulas = [p for p in particulas if p.ativa()]

            # Desenha frutas e partículas
            for fruta in frutas:
                fruta.desenhar(tela)
            for particula in particulas:
                particula.desenhar(tela)

            # Exibe HUD com tempo e pontuação
            desenhar_texto(f"Tempo: {int(tempo_restante)}s", 32, PRETO, 100, 30, centralizado=False)
            desenhar_texto(f"Pontuação: {pontuacao}", 32, PRETO, 700, 30, centralizado=False)

        # Desenha cursor e atualiza tela
        tela.blit(cursor_img, (mouse_x - 32, mouse_y - 32))
        pygame.display.flip()
        relogio.tick(FPS)

    # Encerra pygame
    pygame.quit()


if __name__ == "__main__":
    main()
